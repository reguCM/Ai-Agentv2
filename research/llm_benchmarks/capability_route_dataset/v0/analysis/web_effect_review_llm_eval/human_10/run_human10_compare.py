"""
人間レビュー済み10件 vs Qwen3:8B 一致検証。

- 旧LLM評価11件とは分離（human_10/）
- 人間評価の推測補完は禁止（一覧シートの記入のみ使用）
- Agent / heuristic / search / Gate / Pipeline は変更しない
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from tools.system.config import get_llm_profile
from tools.system.llm import chat

HERE = Path(__file__).resolve().parent
PARENT = HERE.parent
ROOT = PARENT.parent
NORMALIZED = ROOT / "web_effect_review" / "normalized" / "review_dataset.json"
HUMAN_XLSX = ROOT / "web_effect_review_human" / "review_worksheet.xlsx"

CASE_IDS = [
    "B05",
    "C02",
    "C03",
    "P02a",
    "P02b",
    "A06",
    "E02",
    "A03",
    "A04",
    "C04",
]

Q1 = ["◎ 必須に近い", "○ 使う価値が高い", "△ どちらでもよい", "× なくてもよい"]
Q2 = [
    "◎ 十分に有用",
    "○ 一部有用",
    "△ 参考程度",
    "× ほぼ役に立たない",
    "？ 内容不足で判断不能",
]
Q3 = ["Webあり", "Webなし", "ほぼ同等", "どちらも不十分"]
Q4_OPTS = [
    "情報が追加された",
    "情報の鮮度が上がった",
    "根拠が増えた",
    "具体性が増えた",
    "正確性が上がった可能性がある",
    "回答構成だけ変化した",
    "ほぼ変化なし",
    "検索結果をほぼ利用できていない",
    "検索結果が不足していた",
    "その他",
]
Q5 = ["Webなし", "Webあり", "条件付きでWebあり", "条件付きでWebなし"]

# 人間スキーマ → 比較用正規化
HUMAN_Q2_NORM = {
    "◎ とても役立つ": "◎",
    "○ 役立つ": "○",
    "△ 一部役立つ": "△",
    "× 役立たない": "×",
    "― 検索失敗・比較不能": "？",
}
LLM_Q2_NORM = {
    "◎ 十分に有用": "◎",
    "○ 一部有用": "○",
    "△ 参考程度": "△",
    "× ほぼ役に立たない": "×",
    "？ 内容不足で判断不能": "？",
}
HUMAN_CHANGE_NORM = {
    "◎ 明確に良くなった": "改善大",
    "○ 少し良くなった": "改善小",
    "→ ほとんど変わらない": "変化なし",
    "△ 少し悪くなった": "悪化小",
    "× 明確に悪くなった": "悪化大",
    "― 比較不能": "比較不能",
}
ORDER_POS = {"◎": 4, "○": 3, "△": 2, "×": 1, "？": 0, "―": 0}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _msg(response) -> str:
    msg = getattr(response, "message", None)
    if msg is None:
        return ""
    c = getattr(msg, "content", None)
    return "" if c is None else str(c)


def extract_json(text: str) -> dict | None:
    text = (text or "").strip()
    if not text:
        return None
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass
    m = re.search(r"\{[\s\S]*\}", text)
    if not m:
        return None
    try:
        return json.loads(m.group(0))
    except json.JSONDecodeError:
        return None


def validate_choice(value: str, allowed: list[str]) -> str:
    v = str(value or "").strip()
    if v in allowed:
        return v
    for a in allowed:
        if v and (v in a or a.startswith(v[:1])):
            return a
    return v


def load_human_reviews() -> dict[str, dict]:
    if not HUMAN_XLSX.is_file():
        raise FileNotFoundError(f"DATA_MAPPING_ERROR: missing {HUMAN_XLSX}")
    wb = load_workbook(HUMAN_XLSX, data_only=True)
    ws = wb["一覧"]
    headers = [c.value for c in ws[1]]
    by_id: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = {headers[i]: ("" if row[i] is None else row[i]) for i in range(len(headers))}
        cid = str(d.get("case_id") or "").strip()
        if cid in CASE_IDS:
            # 記入済み判定：主要評価列のいずれか
            filled = any(
                str(d.get(k) or "").strip()
                for k in (
                    "Webなし十分性",
                    "検索結果の有用性",
                    "Webあり十分性",
                    "Web検索による変化",
                    "人間メモ",
                )
            )
            if filled:
                by_id[cid] = d
    return by_id


def load_normalized() -> dict[str, dict]:
    data = json.loads(NORMALIZED.read_text(encoding="utf-8"))
    return {c["case_id"]: c for c in data.get("cases") or []}


def build_search_card(case: dict) -> str:
    lines = [
        f"検索クエリ: {case.get('search_query') or ''}",
        f"検索結果件数: {case.get('search_result_count')}",
        f"本文あり件数: {case.get('available_count')} / 本文なし件数: {case.get('empty_count')}",
        f"日本語結果あり: {'はい' if case.get('has_japanese_result') else 'いいえ'}",
        f"英語結果あり: {'はい' if case.get('has_english_result') else 'いいえ'}",
        "",
    ]
    hits = case.get("normalized_hits") or []
    total = len(hits)
    for h in hits:
        idx = h.get("index") or 0
        status = h.get("content_status")
        body = (
            (h.get("original_text_display") or "").strip()
            if status == "available"
            else "【本文なし】"
        )
        if not body:
            body = "【本文なし】"
        summary = (h.get("japanese_summary") or "").strip() or "（要約なし）"
        lang = h.get("result_language") or "unknown"
        lines.extend(
            [
                f"検索結果 {idx} / {total}",
                f"タイトル: {h.get('title') or ''}",
                f"サイト: {h.get('site') or ''}",
                f"URL: {h.get('url') or ''}",
                f"言語: {lang}",
                f"内容取得: {h.get('content_status_label_ja') or status}",
                f"日本語要約: {summary}",
                "原文:",
                body,
                "",
            ]
        )
    return "\n".join(lines).strip()


def build_eval_prompt_card(case: dict, search_card: str) -> str:
    return f"""【ユーザー要求】
{case.get('request') or ''}

【Webなし完成回答】
{case.get('answer_without_web') or ''}

【検索結果】
{search_card}

【Webあり完成回答】
{case.get('answer_with_web') or ''}
""".strip()


EVAL_SYSTEM = """あなたはWeb検索効果のレビュー評価者です。JSONのみ出力。

厳守:
1. 検索結果が空（本文なし）なら、その検索結果自体を有用とは評価しない
2. URLやタイトルだけから本文内容を推測しない
3. Webあり回答が良いからといって検索結果自体を◎にしない
4. 「今」「最新」「最近」「現在」など時間依存性を考慮する
5. Web検索を実行したこと自体を改善とは評価しない
6. WebなしとWebありの内容差を実際に比較する
7. 判断できない場合は無理に◎〜×を選ばず「？」を使う
8. 「Web検索が必要そう」という印象だけで最終価値を◎にしない
9. 検索結果が弱い場合と、Webなし回答が弱い場合を分離する
10. 人間の正解を推測しない
11. 一般知識を検索結果の内容として扱わない
"""


def llm_evaluate(card: str, *, model: str, retries: int = 2) -> dict:
    prompt = f"""以下の材料だけを根拠に評価し、JSONのみ出力せよ。

Q1 Webは必要だったか: {" / ".join(Q1)}
Q2 検索結果は役に立ったか（回答の良さではなく検索結果そのもの）: {" / ".join(Q2)}
Q3 Webなし回答とWebあり回答のどちらが良いか: {" / ".join(Q3)}
Q4 Web検索によって何が変化したか（配列・複数可）: {Q4_OPTS}
Q5 最終的にどちらを採用するか（必ず1つ。どちらも禁止）: {" / ".join(Q5)}
Q6 このケースでは何が起きているか（自由記述・証拠のない断定禁止）
Q7 開発者なら次に何を修正・調査するか（優先最大3個の配列）

出力スキーマ:
{{
  "q1_web_need": "...",
  "q1_reason": "...",
  "q2_search_usefulness": "...",
  "q2_reason": "...",
  "q3_which_better": "...",
  "q3_reason": "...",
  "q4_changes": ["..."],
  "q4_reason": "...",
  "q5_adopt": "...",
  "q5_condition": "条件付きの場合のみ",
  "q5_reason": "...",
  "q6_what_happened": "...",
  "q7_next_fixes": ["1. ...", "2. ...", "3. ..."]
}}

--- 評価材料 ---
{card}
"""
    last = ""
    for attempt in range(retries + 1):
        extra = "\n前回JSON失敗。有効JSONオブジェクト1つのみ。" if attempt else ""
        resp = chat(
            model=model,
            messages=[
                {"role": "system", "content": EVAL_SYSTEM},
                {"role": "user", "content": prompt + extra},
            ],
        )
        last = _msg(resp)
        parsed = extract_json(last)
        if parsed and parsed.get("q2_search_usefulness"):
            parsed["_raw"] = last
            return parsed
    return {"_raw": last}


def normalize_llm(parsed: dict) -> dict:
    q4 = parsed.get("q4_changes") or []
    if isinstance(q4, str):
        q4 = [q4]
    q7 = parsed.get("q7_next_fixes") or []
    if isinstance(q7, str):
        q7 = [q7]
    return {
        "q1_web_need": validate_choice(parsed.get("q1_web_need"), Q1),
        "q1_reason": str(parsed.get("q1_reason") or "").strip(),
        "q2_search_usefulness": validate_choice(parsed.get("q2_search_usefulness"), Q2),
        "q2_reason": str(parsed.get("q2_reason") or "").strip(),
        "q3_which_better": validate_choice(parsed.get("q3_which_better"), Q3),
        "q3_reason": str(parsed.get("q3_reason") or "").strip(),
        "q4_changes": [validate_choice(x, Q4_OPTS) for x in q4 if str(x).strip()],
        "q4_reason": str(parsed.get("q4_reason") or "").strip(),
        "q5_adopt": validate_choice(parsed.get("q5_adopt"), Q5),
        "q5_condition": str(parsed.get("q5_condition") or "").strip(),
        "q5_reason": str(parsed.get("q5_reason") or "").strip(),
        "q6_what_happened": str(parsed.get("q6_what_happened") or "").strip(),
        "q7_next_fixes": [str(x).strip() for x in q7 if str(x).strip()][:3],
    }


def derive_human_comparable(human: dict) -> dict:
    """人間スキーマから比較可能な派生ラベルを作る。未記入は空。推測で埋めない。"""
    search_u = str(human.get("検索結果の有用性") or "").strip()
    without = str(human.get("Webなし十分性") or "").strip()
    with_ = str(human.get("Webあり十分性") or "").strip()
    change = str(human.get("Web検索による変化") or "").strip()
    memo = str(human.get("人間メモ") or "").strip()

    # Q3相当: 明示フィールドから機械的に導出（メモの解釈はしない）
    which = ""
    if change == "× 明確に悪くなった" or (
        without.startswith("◎") and with_.startswith("×")
    ):
        which = "Webなし"
    elif change in ("◎ 明確に良くなった", "○ 少し良くなった") and with_.startswith(
        ("◎", "○")
    ):
        which = "Webあり"
    elif change == "→ ほとんど変わらない" and without.startswith(("◎", "○")) and with_.startswith(
        ("◎", "○")
    ):
        which = "ほぼ同等"
    elif without.startswith("×") and with_.startswith("×"):
        which = "どちらも不十分"
    elif change == "― 比較不能" and without.startswith("×") and with_.startswith("×"):
        which = "どちらも不十分"
    elif change == "― 比較不能" and without.startswith(("◎", "○")):
        which = "Webなし"
    # else leave empty — 導出不能

    # Q5相当
    adopt = ""
    if which == "Webなし":
        adopt = "Webなし"
    elif which == "Webあり":
        adopt = "Webあり"
    elif which == "ほぼ同等" and without.startswith("◎"):
        adopt = "条件付きでWebなし"  # 同等ならなし側維持が自然だが、導出である旨を注記
    elif which == "どちらも不十分":
        adopt = ""  # 強制選択できない

    # Q1相当: 人間は直接記入していない → 空（推測禁止）
    # ただしメモに「WEBなし解答が無理」等がある場合も推測で埋めない

    # Q4相当テキスト
    q4_tags = []
    if change == "→ ほとんど変わらない":
        q4_tags = ["ほぼ変化なし"]
    elif change == "× 明確に悪くなった":
        q4_tags = ["検索結果をほぼ利用できていない", "検索結果が不足していた"]
    elif change == "― 比較不能":
        q4_tags = ["検索結果が不足していた"]
    elif change == "◎ 明確に良くなった":
        q4_tags = ["情報が追加された"]
    elif change == "○ 少し良くなった":
        q4_tags = ["情報が追加された"]

    return {
        "human_question_clarity": str(human.get("質問の明確さ") or "").strip(),
        "human_without_sufficiency": without,
        "human_search_usefulness": search_u,
        "human_with_sufficiency": with_,
        "human_web_change": change,
        "human_memo": memo,
        "derived_q3_which_better": which,
        "derived_q5_adopt": adopt,
        "derived_q4_tags": q4_tags,
        "derived_note": "Q3/Q5は人間スキーマからの機械導出。Q1は人間未記入のため空。",
    }


def choice_agreement(human_norm: str, llm_norm: str) -> str:
    if not human_norm or not llm_norm:
        return ""
    if human_norm == llm_norm:
        return "完全一致"
    ho = ORDER_POS.get(human_norm)
    lo = ORDER_POS.get(llm_norm)
    if ho is not None and lo is not None and abs(ho - lo) == 1:
        return "近似一致"
    return "不一致"


def map_llm_q2_norm(v: str) -> str:
    return LLM_Q2_NORM.get(v, "")


def map_human_q2_norm(v: str) -> str:
    return HUMAN_Q2_NORM.get(v, "")


def map_change_bucket_from_llm_q4(tags: list[str]) -> str:
    s = set(tags or [])
    if "ほぼ変化なし" in s and len(s) == 1:
        return "変化なし"
    if "検索結果が不足していた" in s or "検索結果をほぼ利用できていない" in s:
        if any(
            x in s
            for x in (
                "情報が追加された",
                "情報の鮮度が上がった",
                "根拠が増えた",
                "具体性が増えた",
            )
        ):
            return "混在"
        return "比較不能"  # 近似: 不足・失敗寄り
    if any(
        x in s
        for x in (
            "情報が追加された",
            "情報の鮮度が上がった",
            "根拠が増えた",
            "具体性が増えた",
            "正確性が上がった可能性がある",
        )
    ):
        return "改善小"
    if "回答構成だけ変化した" in s:
        return "変化なし"
    return "不明"


def llm_compare_structure(
    *,
    human_memo: str,
    human_fields: dict,
    llm_q6: str,
    llm_q7: list[str],
    model: str,
) -> dict:
    """レベル2–4: 理由・問題構造・改善方針の一致をLLMに判定（元文も保存）。"""
    prompt = f"""人間レビューとLLM評価の「問題認識」が一致するか判定し、JSONのみ出力。

人間は正解ではない。一致度だけ測る。

人間フィールド:
- Webなし十分性: {human_fields.get('human_without_sufficiency')}
- 検索結果の有用性: {human_fields.get('human_search_usefulness')}
- Webあり十分性: {human_fields.get('human_with_sufficiency')}
- Web検索による変化: {human_fields.get('human_web_change')}
- 人間メモ: {human_memo or '（空）'}

LLM:
- Q6問題構造: {llm_q6}
- Q7次の修正: {llm_q7}

出力:
{{
  "reason_agreement": "同じ問題を指摘 / 一部共通 / 異なる / 判断不能",
  "reason_note": "短い日本語",
  "structure_agreement": "一致 / 部分一致 / 不一致 / 判断不能",
  "structure_note": "短い日本語（同じ現象を見ているか）",
  "fix_agreement": "一致 / 部分一致 / 不一致 / 判断不能",
  "fix_note": "短い日本語（次に直すべき場所が同じか）",
  "conclusion_same_cause_different": true/false,
  "special_note": "結論一致だが原因不一致などの指摘"
}}
"""
    resp = chat(
        model=model,
        messages=[
            {
                "role": "system",
                "content": "比較判定者。JSONのみ。人間を正解扱いにしない。",
            },
            {"role": "user", "content": prompt},
        ],
    )
    raw = _msg(resp)
    parsed = extract_json(raw) or {}
    parsed["_raw"] = raw
    return parsed


def llm_self_check(eval_card: str, llm_eval: dict, *, model: str) -> dict:
    prompt = f"""あなた自身の直前評価を自己検証せよ。JSONのみ。

チェック項目（各 true/false + 短い理由）:
1. 回答の良し悪しと検索結果の有用性を混同していないか
2. 検索結果本文を実際に根拠として評価したか
3. 「Webなし回答が良い」ことを理由に「検索結果が有用」と評価していないか
4. 時間依存情報を古い知識で回答しているケースを見落としていないか
5. 「検索結果がない」ことと「検索結果が不要」だったことを混同していないか

評価材料:
{eval_card}

あなたの評価:
{json.dumps(llm_eval, ensure_ascii=False, indent=2)}

出力:
{{
  "no_confusion_answer_vs_search": true/false,
  "used_search_body_as_evidence": true/false,
  "did_not_infer_search_from_without_web": true/false,
  "noticed_time_sensitivity": true/false,
  "did_not_confuse_empty_search_with_unnecessary": true/false,
  "self_check_notes": "短い日本語"
}}
"""
    resp = chat(
        model=model,
        messages=[
            {"role": "system", "content": "自己検証。JSONのみ。"},
            {"role": "user", "content": prompt},
        ],
    )
    raw = _msg(resp)
    parsed = extract_json(raw) or {}
    parsed["_raw"] = raw
    return parsed


def write_all(rows: list[dict], *, model: str, mapping_errors: list[str]) -> None:
    HERE.mkdir(parents=True, exist_ok=True)

    payload = {
        "kind": "human_10_vs_qwen_eval",
        "ts": _now(),
        "model": model,
        "case_ids": CASE_IDS,
        "case_count": len(rows),
        "mapping_errors": mapping_errors,
        "old_llm_eval_11_not_used": True,
        "human_not_treated_as_ground_truth": True,
        "not_connected_to_agent": True,
        "evaluations": rows,
    }
    (HERE / "llm_evaluation.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    # llm_evaluation.csv
    fields = [
        "case_id",
        "observation_id",
        "request",
        "q1_web_need",
        "q1_reason",
        "q2_search_usefulness",
        "q2_reason",
        "q3_which_better",
        "q3_reason",
        "q4_changes",
        "q5_adopt",
        "q5_reason",
        "q6_what_happened",
        "q7_next_fixes",
    ]
    with (HERE / "llm_evaluation.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            llm = r["llm"]
            w.writerow(
                {
                    "case_id": r["case_id"],
                    "observation_id": r["observation_id"],
                    "request": r["request"],
                    "q1_web_need": llm["q1_web_need"],
                    "q1_reason": llm["q1_reason"],
                    "q2_search_usefulness": llm["q2_search_usefulness"],
                    "q2_reason": llm["q2_reason"],
                    "q3_which_better": llm["q3_which_better"],
                    "q3_reason": llm["q3_reason"],
                    "q4_changes": ";".join(llm["q4_changes"]),
                    "q5_adopt": llm["q5_adopt"],
                    "q5_reason": llm["q5_reason"],
                    "q6_what_happened": llm["q6_what_happened"],
                    "q7_next_fixes": " | ".join(llm["q7_next_fixes"]),
                }
            )

    # human_vs_llm.csv
    cmp_fields = [
        "case_id",
        "human_search_usefulness",
        "llm_q2",
        "q2_agreement",
        "human_web_change",
        "llm_q4_bucket",
        "q4_agreement",
        "human_derived_q3",
        "llm_q3",
        "q3_agreement",
        "human_derived_q5",
        "llm_q5",
        "q5_agreement",
        "human_without",
        "human_with",
        "reason_agreement",
        "structure_agreement",
        "fix_agreement",
        "conclusion_same_cause_different",
        "human_memo",
        "llm_q6",
        "llm_q7",
    ]
    with (HERE / "human_vs_llm.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cmp_fields)
        w.writeheader()
        for r in rows:
            c = r["comparison"]
            h = r["human_comparable"]
            llm = r["llm"]
            w.writerow(
                {
                    "case_id": r["case_id"],
                    "human_search_usefulness": h["human_search_usefulness"],
                    "llm_q2": llm["q2_search_usefulness"],
                    "q2_agreement": c.get("q2_agreement", ""),
                    "human_web_change": h["human_web_change"],
                    "llm_q4_bucket": c.get("llm_q4_bucket", ""),
                    "q4_agreement": c.get("q4_agreement", ""),
                    "human_derived_q3": h["derived_q3_which_better"],
                    "llm_q3": llm["q3_which_better"],
                    "q3_agreement": c.get("q3_agreement", ""),
                    "human_derived_q5": h["derived_q5_adopt"],
                    "llm_q5": llm["q5_adopt"],
                    "q5_agreement": c.get("q5_agreement", ""),
                    "human_without": h["human_without_sufficiency"],
                    "human_with": h["human_with_sufficiency"],
                    "reason_agreement": c.get("reason_agreement", ""),
                    "structure_agreement": c.get("structure_agreement", ""),
                    "fix_agreement": c.get("fix_agreement", ""),
                    "conclusion_same_cause_different": c.get(
                        "conclusion_same_cause_different", ""
                    ),
                    "human_memo": h["human_memo"],
                    "llm_q6": llm["q6_what_happened"],
                    "llm_q7": " | ".join(llm["q7_next_fixes"]),
                }
            )

    # evaluation_cards.md
    md = ["# 人間10件 vs Qwen3:8B 評価カード", ""]
    for r in rows:
        h = r["human_comparable"]
        llm = r["llm"]
        c = r["comparison"]
        md.extend(
            [
                "━━━━━━━━━━━━━━━━━━━━",
                f"ケース：{r['case_id']}",
                "",
                "【ユーザー要求】",
                r["request"],
                "",
                "【Webなし完成回答】",
                r["answer_without_web"],
                "",
                "【検索結果】",
                r["search_card"],
                "",
                "【Webあり完成回答】",
                r["answer_with_web"],
                "",
                "【Qwen評価】",
                f"Q1 Web必要性：{llm['q1_web_need']} — {llm['q1_reason']}",
                f"Q2 検索結果の有用性：{llm['q2_search_usefulness']} — {llm['q2_reason']}",
                f"Q3 どちらの回答：{llm['q3_which_better']} — {llm['q3_reason']}",
                f"Q4 Webで何が変化：{'; '.join(llm['q4_changes'])} — {llm['q4_reason']}",
                f"Q5 最終採用：{llm['q5_adopt']} — {llm['q5_reason']}",
                "",
                "【Qwenが認識した問題】",
                llm["q6_what_happened"],
                "",
                "【Qwenが提案する次の修正】",
                *llm["q7_next_fixes"],
                "",
                "【人間レビュー】",
                f"質問の明確さ：{h['human_question_clarity']}",
                f"Webなし十分性：{h['human_without_sufficiency']}",
                f"検索結果の有用性：{h['human_search_usefulness']}",
                f"Webあり十分性：{h['human_with_sufficiency']}",
                f"Web検索による変化：{h['human_web_change']}",
                f"人間メモ：{h['human_memo'] or '（空）'}",
                f"（導出Q3：{h['derived_q3_which_better'] or '導出不能'} / 導出Q5：{h['derived_q5_adopt'] or '導出不能'}）",
                "",
                "【一致状況】",
                f"個別評価 Q2：{c.get('q2_agreement','')}",
                f"個別評価 Q3：{c.get('q3_agreement','')}",
                f"個別評価 Q4：{c.get('q4_agreement','')}",
                f"個別評価 Q5：{c.get('q5_agreement','')}",
                f"理由：{c.get('reason_agreement','')} — {c.get('reason_note','')}",
                f"問題構造：{c.get('structure_agreement','')} — {c.get('structure_note','')}",
                f"改善方針：{c.get('fix_agreement','')} — {c.get('fix_note','')}",
                "━━━━━━━━━━━━━━━━━━━━",
                "",
            ]
        )
    (HERE / "evaluation_cards.md").write_text("\n".join(md), encoding="utf-8")

    # SELF_CHECK.md
    sc = ["# SELF_CHECK（評価器自己検証）", ""]
    for r in rows:
        s = r.get("self_check") or {}
        sc.extend(
            [
                f"## {r['case_id']}",
                f"- 回答と検索の混同なし: {s.get('no_confusion_answer_vs_search')}",
                f"- 本文を根拠にした: {s.get('used_search_body_as_evidence')}",
                f"- Webなし良さ→検索有用の誤推論なし: {s.get('did_not_infer_search_from_without_web')}",
                f"- 時間依存を認識: {s.get('noticed_time_sensitivity')}",
                f"- 空検索≠不要の混同なし: {s.get('did_not_confuse_empty_search_with_unnecessary')}",
                f"- メモ: {s.get('self_check_notes','')}",
                "",
            ]
        )
    (HERE / "SELF_CHECK.md").write_text("\n".join(sc), encoding="utf-8")

    write_report(rows, model=model, mapping_errors=mapping_errors)

    (HERE / "README.md").write_text(
        "\n".join(
            [
                "# 人間レビュー10件 vs Qwen3:8B",
                "",
                "対象: B05, C02, C03, P02a, P02b, A06, E02, A03, A04, C04",
                "",
                "- 旧LLM評価11件（WB系含む）は**比較に使用していない**",
                "- 人間評価を正解ラベルとして扱わない（一致度の測定）",
                "- 人間未記入のQ1は空欄（推測禁止）",
                "- Q3/Q5人間側はスキーマ差分のため機械導出（`derived_*`）",
                "",
                "## 実行",
                "",
                "```text",
                "python run_human10_compare.py",
                "```",
                "",
            ]
        ),
        encoding="utf-8",
    )


def write_report(rows: list[dict], *, model: str, mapping_errors: list[str]) -> None:
    def rate(key: str, good: set[str]) -> str:
        vals = [r["comparison"].get(key, "") for r in rows]
        filled = [v for v in vals if v]
        if not filled:
            return "（比較不能・未算出）"
        ok = sum(1 for v in filled if v in good)
        return f"{ok}/{len(filled)} = {100*ok/len(filled):.0f}%（分母=比較可能な件数）"

    q2_exact = rate("q2_agreement", {"完全一致"})
    q2_near = rate("q2_agreement", {"完全一致", "近似一致"})
    q3 = rate("q3_agreement", {"完全一致"})
    q4 = rate("q4_agreement", {"完全一致", "近似一致"})
    q5 = rate("q5_agreement", {"完全一致", "近似一致"})
    struct = rate("structure_agreement", {"一致", "部分一致"})
    struct_exact = rate("structure_agreement", {"一致"})
    fix = rate("fix_agreement", {"一致", "部分一致"})
    fix_exact = rate("fix_agreement", {"一致"})
    reason = rate("reason_agreement", {"同じ問題を指摘", "一部共通"})

    empty_search_cases = [
        r["case_id"]
        for r in rows
        if r.get("available_count", 0) == 0
    ]
    time_cases = [
        r
        for r in rows
        if any(
            w in (r["request"] or "")
            for w in ("今", "最近", "最新", "現在")
        )
    ]

    notable = []
    for r in rows:
        c = r["comparison"]
        if c.get("conclusion_same_cause_different"):
            notable.append(
                f"- `{r['case_id']}`: 結論一致だが原因不一致の可能性 — {c.get('special_note','')}"
            )
        if c.get("structure_agreement") == "不一致":
            notable.append(
                f"- `{r['case_id']}`: 問題構造不一致 — H:{r['human_comparable']['human_memo'][:80]} / L:{r['llm']['q6_what_happened'][:80]}"
            )

    lines = [
        "# REPORT — 人間10件 vs Qwen3:8B",
        "",
        f"- モデル: `{model}`",
        f"- 件数: **{len(rows)}**（指定10件のみ）",
        f"- マッピングエラー: {mapping_errors or 'なし'}",
        "- 旧評価11件は未使用",
        "- 人間を正解ラベルとして扱っていない",
        "- サンプル少数（n=10）。統計的断定はしない",
        "",
        "## A. 個別評価一致率（レベル1）",
        "",
        f"- Q2 検索有用性 完全一致: {q2_exact}",
        f"- Q2 完全+近似: {q2_near}",
        f"- Q3 どちらが良いか: {q3}",
        f"- Q4 変化（バケット）: {q4}",
        f"- Q5 採用（導出人間 vs LLM）: {q5}",
        "- Q1 Web必要性: 人間スキーマに直接項目なし → **未算出**",
        "",
        "## B. 理由一致率（レベル2）",
        "",
        f"- {reason}",
        "",
        "## C. 問題構造一致率（レベル3）※重視",
        "",
        f"- 一致のみ: {struct_exact}",
        f"- 一致+部分一致: {struct}",
        "",
        "## D. 改善方針一致率（レベル4）※重視",
        "",
        f"- 一致のみ: {fix_exact}",
        f"- 一致+部分一致: {fix}",
        "",
        "## 特記ケース",
        "",
        *(notable or ["（特記なし）"]),
        "",
        f"## 本文なしケース: {empty_search_cases}",
        f"## 時間依存語を含むケース: {[r['case_id'] for r in time_cases]}",
        "",
        "## 最終8問への回答",
        "",
        "### 1. 個別評価の一致",
        f"Q2（検索有用性）完全一致 {q2_exact}、近似込み {q2_near}。Q3/Q5は人間スキーマ差分のため導出比較。詳細は `human_vs_llm.csv`。",
        "",
        "### 2. 検索結果そのものの有用性を評価できたか",
        "本文なし多数の本セットでは、Q2で「？/×」側に寄せられるかが要点。各ケースのQ2と人間「検索失敗・比較不能」の対応を参照。",
        "",
        "### 3. 「Webなしでも十分」と「検索機能が弱い」の区別",
        "問題構造一致（C）と特記ケースを参照。区別できているかはQ6と人間メモの対応が指標。",
        "",
        "### 4. 時間依存性の認識",
        f"時間依存語ケース { [r['case_id'] for r in time_cases] } のQ1/Q6とSELF_CHECKの `noticed_time_sensitivity` を参照。",
        "",
        "### 5. 検索結果の問題を原因として認識できたか",
        "レベル3（問題構造）一致率が主指標。本文取得失敗を「Web不要」と混同していないかも確認。",
        "",
        "### 6. 「次に何を修正すべきか」の一致",
        f"レベル4: {fix}（一致のみ {fix_exact}）。ここが低い場合、セル一致が高くても評価器委任は尚早。",
        "",
        "### 7. 90%以上で任せられるか",
        "**n=10の予備検証であり、90%超でも十分とは断定しない。** C/D（問題構造・改善方針）が安定して高い場合に限り、一次評価器候補として追加検証（20–30件）を検討する段階。",
        "",
        "### 8. まだ人間レビューが必要な領域",
        "- 質問意味の破綻・曖昧さ（例: C03）",
        "- 時間依存で検索失敗時の「必要だが結果が空」判断",
        "- 結論は同じだが原因が違うケース",
        "- Q7の具体的修正対象の優先順位",
        "",
        "## 次の判断（自動結論禁止）",
        "",
        "人間が A/B/C/D（追加検証 / プロンプト改善 / 評価器修正 / 基準見直し）を選択する。",
        "本レポートは検索機能・Stage3/4実装判断を自動では行わない。",
        "",
    ]
    (HERE / "REPORT.md").write_text("\n".join(lines), encoding="utf-8")


def build_comparison(human_c: dict, llm: dict, struct: dict) -> dict:
    h_q2 = map_human_q2_norm(human_c["human_search_usefulness"])
    l_q2 = map_llm_q2_norm(llm["q2_search_usefulness"])
    q2_agr = choice_agreement(h_q2, l_q2)

    h_ch = HUMAN_CHANGE_NORM.get(human_c["human_web_change"], "")
    l_ch = map_change_bucket_from_llm_q4(llm["q4_changes"])
    if not h_ch or l_ch in ("不明", "混在"):
        q4_agr = ""
    elif h_ch == l_ch:
        q4_agr = "完全一致"
    elif {h_ch, l_ch} <= {"改善大", "改善小"} or {h_ch, l_ch} <= {"悪化大", "悪化小"}:
        q4_agr = "近似一致"
    elif h_ch == "比較不能" and l_ch == "比較不能":
        q4_agr = "完全一致"
    elif h_ch == "比較不能" and l_ch in ("変化なし", "比較不能"):
        q4_agr = "近似一致"
    else:
        q4_agr = "不一致"

    h_q3 = human_c["derived_q3_which_better"]
    l_q3 = llm["q3_which_better"]
    if not h_q3:
        q3_agr = ""
    elif h_q3 == l_q3:
        q3_agr = "完全一致"
    elif {h_q3, l_q3} <= {"ほぼ同等", "Webなし"} and human_c[
        "human_web_change"
    ] == "→ ほとんど変わらない":
        q3_agr = "近似一致"
    else:
        q3_agr = "不一致"

    h_q5 = human_c["derived_q5_adopt"]
    l_q5 = llm["q5_adopt"]
    if not h_q5:
        q5_agr = ""
    elif h_q5 == l_q5:
        q5_agr = "完全一致"
    elif h_q5.replace("条件付きで", "") == l_q5.replace("条件付きで", ""):
        q5_agr = "近似一致"
    else:
        q5_agr = "不一致"

    def _norm_reason(v: str) -> str:
        v = str(v or "").strip()
        alias = {
            "一致": "同じ問題を指摘",
            "部分一致": "一部共通",
            "不一致": "異なる",
        }
        return alias.get(v, v)

    def _norm_struct(v: str) -> str:
        v = str(v or "").strip()
        alias = {"同じ問題を指摘": "一致", "一部共通": "部分一致", "異なる": "不一致"}
        return alias.get(v, v)

    return {
        "q2_agreement": q2_agr,
        "q4_agreement": q4_agr,
        "q3_agreement": q3_agr,
        "q5_agreement": q5_agr,
        "llm_q4_bucket": l_ch,
        "reason_agreement": _norm_reason(struct.get("reason_agreement", "")),
        "reason_note": struct.get("reason_note", ""),
        "structure_agreement": _norm_struct(struct.get("structure_agreement", "")),
        "structure_note": struct.get("structure_note", ""),
        "fix_agreement": _norm_struct(struct.get("fix_agreement", "")),
        "fix_note": struct.get("fix_note", ""),
        "conclusion_same_cause_different": bool(
            struct.get("conclusion_same_cause_different")
        ),
        "special_note": struct.get("special_note", ""),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--case-id", action="append", default=[])
    args = parser.parse_args(argv)

    model = get_llm_profile()["model"]
    target_ids = args.case_id or list(CASE_IDS)
    humans = load_human_reviews()
    normalized = load_normalized()

    mapping_errors: list[str] = []
    rows: list[dict] = []

    for cid in target_ids:
        if cid not in CASE_IDS:
            mapping_errors.append(f"DATA_MAPPING_ERROR: {cid} not in human-10 list")
            continue
        if cid not in humans:
            mapping_errors.append(f"DATA_MAPPING_ERROR: human review missing for {cid}")
            continue
        if cid not in normalized:
            mapping_errors.append(
                f"DATA_MAPPING_ERROR: normalized case missing for {cid}"
            )
            continue
        case = normalized[cid]
        human = humans[cid]
        # observation_id 一致確認
        h_obs = str(human.get("observation_id") or "").strip()
        c_obs = str(case.get("observation_id") or "").strip()
        if h_obs and c_obs and h_obs != c_obs:
            mapping_errors.append(
                f"DATA_MAPPING_ERROR: observation_id mismatch {cid}: {h_obs} vs {c_obs}"
            )
            continue

        search_card = build_search_card(case)
        eval_card = build_eval_prompt_card(case, search_card)
        human_c = derive_human_comparable(human)

        if args.dry_run:
            llm = {
                "q1_web_need": "",
                "q1_reason": "dry-run",
                "q2_search_usefulness": "",
                "q2_reason": "",
                "q3_which_better": "",
                "q3_reason": "",
                "q4_changes": [],
                "q4_reason": "",
                "q5_adopt": "",
                "q5_condition": "",
                "q5_reason": "",
                "q6_what_happened": "",
                "q7_next_fixes": [],
            }
            struct = {}
            self_check = {}
        else:
            print(f"[{cid}] LLM evaluate...")
            llm = normalize_llm(llm_evaluate(eval_card, model=model))
            print(f"[{cid}] structure compare...")
            struct = llm_compare_structure(
                human_memo=human_c["human_memo"],
                human_fields=human_c,
                llm_q6=llm["q6_what_happened"],
                llm_q7=llm["q7_next_fixes"],
                model=model,
            )
            print(f"[{cid}] self-check...")
            self_check = llm_self_check(eval_card, llm, model=model)

        comparison = build_comparison(human_c, llm, struct)
        rows.append(
            {
                "case_id": cid,
                "observation_id": c_obs or h_obs,
                "request": case.get("request") or "",
                "answer_without_web": case.get("answer_without_web") or "",
                "answer_with_web": case.get("answer_with_web") or "",
                "search_query": case.get("search_query") or "",
                "available_count": case.get("available_count"),
                "empty_count": case.get("empty_count"),
                "search_card": search_card,
                "human_raw": {
                    k: human.get(k)
                    for k in (
                        "質問の明確さ",
                        "Webなし十分性",
                        "検索結果の有用性",
                        "Webあり十分性",
                        "Web検索による変化",
                        "人間メモ",
                    )
                },
                "human_comparable": human_c,
                "llm": llm,
                "comparison": comparison,
                "self_check": self_check,
            }
        )

    write_all(rows, model=model, mapping_errors=mapping_errors)
    print(f"done n={len(rows)} errors={mapping_errors} -> {HERE}")
    return 1 if mapping_errors and not rows else 0


if __name__ == "__main__":
    raise SystemExit(main())
