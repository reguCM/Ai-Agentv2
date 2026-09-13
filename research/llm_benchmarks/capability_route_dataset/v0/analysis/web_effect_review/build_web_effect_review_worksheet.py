"""
既存 hits>0 の25件を、人間レビュー用シートに再構成する。

- レビュー資料のみ。heuristic / Gate / Pipeline / Stage3/4 は変更しない
- 既存 JSON/生データは削除しない
- 人間ラベルは空欄のまま（レビュー結果は作らない）
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "web_effect_review_hits.json"

QUESTION_CLARITY = ["◎ 明確", "○ ほぼ明確", "△ 曖昧", "× 意味が破綻している"]
WITHOUT_QUALITY = ["◎ 十分", "○ 概ね十分", "△ 情報不足", "× 明確な問題あり"]
SEARCH_USEFUL = ["◎ 非常に有用", "○ 有用", "△ 一部有用", "× 無関係・取得失敗"]
WEB_EFFECT = ["◎ 明確に改善", "○ 少し改善", "△ ほぼ変化なし", "▲ 悪化", "× 明確に悪化"]
UTILIZATION = ["◎ 十分活用", "○ 一部活用", "△ ほぼ活用されていない", "× 活用されていない", "― 判断不能"]


def format_search_results_text(web_search: dict | None) -> str:
    ws = web_search or {}
    res = ws.get("result") or {}
    hits = res.get("hits") or []
    lines = [
        f"query: {ws.get('query') or res.get('query') or ''}",
        f"hit_count: {res.get('hit_count') if res.get('hit_count') is not None else len(hits)}",
        f"error: {res.get('error') or '(none)'}",
        f"backends_tried: {res.get('backends_tried')}",
        "",
    ]
    if not hits:
        lines.append("（検索結果なし / 取得失敗）")
        return "\n".join(lines)
    for i, hit in enumerate(hits, 1):
        if not isinstance(hit, dict):
            lines.append(f"[{i}] {hit!r}")
            continue
        lines.append(f"[{i}] title: {hit.get('title') or ''}")
        lines.append(f"    url: {hit.get('url') or ''}")
        snippet = hit.get("snippet") or hit.get("description") or ""
        lines.append(f"    snippet: {snippet}")
        lines.append("")
    return "\n".join(lines).rstrip()


def load_rows() -> list[dict]:
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    rows: list[dict] = []
    for case in data.get("cases") or []:
        ws = case.get("web_search") or {}
        res = ws.get("result") or {}
        hits = res.get("hits") if isinstance(res.get("hits"), list) else []
        rows.append(
            {
                "case_id": case.get("case_id"),
                "request": case.get("request") or "",
                "category": case.get("band") or "",
                "answer_without_web": case.get("answer_without_web") or "",
                "search_query": ws.get("query") or res.get("query") or "",
                "search_result_count": res.get("hit_count")
                if res.get("hit_count") is not None
                else len(hits),
                "search_results": hits,
                "search_results_text": format_search_results_text(ws),
                "search_results_json": json.dumps(hits, ensure_ascii=False, indent=2),
                "search_error": res.get("error"),
                "answer_with_web": case.get("answer_with_web") or "",
                "observation_id": None,
                "source_file": "web_effect_review_hits.json",
                # 人間記入（空）
                "question_clarity": "",
                "web_without_answer_quality": "",
                "search_result_usefulness": "",
                "web_effect": "",
                "search_result_utilization": "",
                "human_reason": "",
            }
        )
    return rows


def write_csv(rows: list[dict], path: Path) -> None:
    fields = [
        "case_id",
        "request",
        "category",
        "question_clarity",
        "answer_without_web",
        "search_query",
        "search_result_count",
        "search_results",
        "answer_with_web",
        "web_without_answer_quality",
        "search_result_usefulness",
        "web_effect",
        "search_result_utilization",
        "human_reason",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            out = dict(row)
            out["search_results"] = row["search_results_json"]
            writer.writerow(out)


def write_json(rows: list[dict], path: Path) -> None:
    payload = {
        "kind": "web_effect_human_review_worksheet",
        "source": str(SOURCE.name),
        "not_agent_gold": True,
        "no_auto_labels": True,
        "review_restructure_only": True,
        "case_count": len(rows),
        "counts": {
            "answer_without_web_nonempty": sum(
                1 for r in rows if (r.get("answer_without_web") or "").strip()
            ),
            "search_results_present": sum(
                1 for r in rows if (r.get("search_result_count") or 0) > 0
            ),
            "answer_with_web_nonempty": sum(
                1 for r in rows if (r.get("answer_with_web") or "").strip()
            ),
        },
        "dropdowns": {
            "question_clarity": QUESTION_CLARITY,
            "web_without_answer_quality": WITHOUT_QUALITY,
            "search_result_usefulness": SEARCH_USEFUL,
            "web_effect": WEB_EFFECT,
            "search_result_utilization": UTILIZATION,
        },
        "cases": rows,
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_compare_md(rows: list[dict], path: Path) -> None:
    lines = [
        "# WEB効果レビュー比較カード（hits>0・25件）",
        "",
        "評価は「どちらを採用するか」ではない。差分観測用。",
        "選択肢の意味は `README.md` / Excelの Choices シートを参照。",
        "",
    ]
    for row in rows:
        lines.extend(
            [
                f"## {row['case_id']} · category={row.get('category')}",
                "",
                "### ① ユーザー要求",
                "",
                "```text",
                row.get("request") or "(empty)",
                "```",
                "",
                "> この質問は意味が明確ですか？（Excel: question_clarity）",
                "",
                "### ② Webなし完成回答（answer_without_web）",
                "",
                "```text",
                row.get("answer_without_web") or "(empty)",
                "```",
                "",
                "> Webなし回答だけで十分ですか？（Excel: web_without_answer_quality）",
                "",
                "### ③ Web検索結果",
                "",
                f"- query: `{row.get('search_query')}`",
                f"- result_count: `{row.get('search_result_count')}`",
                "",
                "```text",
                row.get("search_results_text") or "(empty)",
                "```",
                "",
                "```json",
                row.get("search_results_json") or "[]",
                "```",
                "",
                "> 検索結果は回答に使える内容でしたか？（Excel: search_result_usefulness）",
                "",
                "### ④ Webあり完成回答（answer_with_web）",
                "",
                "```text",
                row.get("answer_with_web") or "(empty)",
                "```",
                "",
                "> Web検索によって完成回答はどう変化しましたか？（Excel: web_effect）",
                "",
                "> 検索結果は最終回答に活用されましたか？（Excel: search_result_utilization）",
                "",
                "> 開発上のメモ（Excel: human_reason）",
                "",
                "---",
                "",
            ]
        )
    path.write_text("\n".join(lines), encoding="utf-8")


def _add_dv(ws, col_letter: str, choices: list[str], max_row: int) -> None:
    formula = '"' + ",".join(choices) + '"'
    dv = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True,
        showDropDown=False,
    )
    dv.error = "リストから選択してください"
    dv.errorTitle = "入力エラー"
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{max_row}")


def write_xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()

    # --- Choices sheet ---
    ws_ch = wb.active
    ws_ch.title = "Choices"
    ws_ch["A1"] = "field"
    ws_ch["B1"] = "option"
    ws_ch["C1"] = "meaning"
    meanings = {
        "question_clarity": {
            "◎ 明確": "要求の意図がはっきりしている",
            "○ ほぼ明確": "おおむね通るが細部は曖昧",
            "△ 曖昧": "解釈が複数ありうる",
            "× 意味が破綻している": "レビュー上、質問自体が成立しにくい（元requestは変更しない）",
        },
        "web_without_answer_quality": {
            "◎ 十分": "Webなし完成回答だけでユーザーに返してよい",
            "○ 概ね十分": "大きな欠落はないが改善余地あり",
            "△ 情報不足": "不足が目立つ",
            "× 明確な問題あり": "誤り・空疎・破綻など明確な問題",
        },
        "search_result_usefulness": {
            "◎ 非常に有用": "回答に直結する良質な結果",
            "○ 有用": "使える",
            "△ 一部有用": "一部だけ使える",
            "× 無関係・取得失敗": "使えない、または取得失敗",
        },
        "web_effect": {
            "◎ 明確に改善": "Webあり完成回答が明確に良くなった",
            "○ 少し改善": "改善はあるが小さい",
            "△ ほぼ変化なし": "実質差が小さい",
            "▲ 悪化": "やや悪化",
            "× 明確に悪化": "明確に悪化",
        },
        "search_result_utilization": {
            "◎ 十分活用": "検索結果が最終回答に十分反映",
            "○ 一部活用": "一部反映",
            "△ ほぼ活用されていない": "ほぼ未反映",
            "× 活用されていない": "未反映",
            "― 判断不能": "判断材料不足",
        },
    }
    r = 2
    for field, options in (
        ("question_clarity", QUESTION_CLARITY),
        ("web_without_answer_quality", WITHOUT_QUALITY),
        ("search_result_usefulness", SEARCH_USEFUL),
        ("web_effect", WEB_EFFECT),
        ("search_result_utilization", UTILIZATION),
    ):
        for opt in options:
            ws_ch.cell(r, 1, field)
            ws_ch.cell(r, 2, opt)
            ws_ch.cell(r, 3, meanings[field].get(opt, ""))
            r += 1

    # --- Overview sheet (wide, for filtering) ---
    ws = wb.create_sheet("Overview", 0)
    headers = [
        "case_id",
        "category",
        "request",
        "question_clarity",
        "answer_without_web",
        "web_without_answer_quality",
        "search_query",
        "search_result_count",
        "search_results",
        "search_result_usefulness",
        "answer_with_web",
        "web_effect",
        "search_result_utilization",
        "human_reason",
    ]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for col, name in enumerate(headers, 1):
        cell = ws.cell(1, col, name)
        cell.fill = header_fill
        cell.font = header_font
    wrap = Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(rows, 2):
        values = [
            row["case_id"],
            row["category"],
            row["request"],
            "",
            row["answer_without_web"],
            "",
            row["search_query"],
            row["search_result_count"],
            row["search_results_text"],
            "",
            row["answer_with_web"],
            "",
            "",
            "",
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(i, col, val)
            cell.alignment = wrap
        ws.row_dimensions[i].height = 120

    max_row = len(rows) + 1
    # dropdown columns: D=question_clarity, F=web_without, J=search_useful, L=web_effect, M=utilization
    _add_dv(ws, "D", QUESTION_CLARITY, max_row)
    _add_dv(ws, "F", WITHOUT_QUALITY, max_row)
    _add_dv(ws, "J", SEARCH_USEFUL, max_row)
    _add_dv(ws, "L", WEB_EFFECT, max_row)
    _add_dv(ws, "M", UTILIZATION, max_row)

    widths = {
        "A": 10,
        "B": 8,
        "C": 28,
        "D": 18,
        "E": 42,
        "F": 18,
        "G": 22,
        "H": 10,
        "I": 42,
        "J": 18,
        "K": 42,
        "L": 16,
        "M": 20,
        "N": 28,
    }
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:N{max_row}"

    # --- One card sheet per case (compact review form) ---
    label_fill = PatternFill("solid", fgColor="D6EAF8")
    for row in rows:
        # sheet name max 31 chars
        title = str(row["case_id"])[:31]
        card = wb.create_sheet(title)
        card["A1"] = "項目"
        card["B1"] = "内容"
        card["A1"].font = header_font
        card["B1"].font = header_font
        card["A1"].fill = header_fill
        card["B1"].fill = header_fill

        blocks = [
            ("case_id", row["case_id"]),
            ("category", row["category"]),
            ("① ユーザー要求", row["request"]),
            ("質問の明確さ（記入）", ""),
            ("② Webなし完成回答", row["answer_without_web"]),
            ("Webなし回答の十分性（記入）", ""),
            ("検索クエリ", row["search_query"]),
            ("検索結果件数", row["search_result_count"]),
            ("③ Web検索結果（本文）", row["search_results_text"]),
            ("③ Web検索結果（JSON生）", row["search_results_json"]),
            ("検索結果の有用性（記入）", ""),
            ("④ Webあり完成回答", row["answer_with_web"]),
            ("Webによる変化（記入）", ""),
            ("検索結果の活用（記入）", ""),
            ("開発上のメモ（記入）", ""),
        ]
        for i, (label, value) in enumerate(blocks, 2):
            a = card.cell(i, 1, label)
            b = card.cell(i, 2, value)
            a.fill = label_fill
            a.alignment = Alignment(wrap_text=True, vertical="top")
            b.alignment = Alignment(wrap_text=True, vertical="top")
            if "記入" in label:
                card.row_dimensions[i].height = 30
            else:
                card.row_dimensions[i].height = 90 if len(str(value)) > 80 else 45

        # dropdowns on fill rows: 5,7,12,14,15 (1-index in blocks -> excel row = index+1)
        # blocks rows: 2..16 -> fill at excel rows:
        # 質問の明確さ = row 5, Webなし十分性=7, 検索有用=12, Web変化=14, 活用=15
        for excel_row, choices in (
            (5, QUESTION_CLARITY),
            (7, WITHOUT_QUALITY),
            (12, SEARCH_USEFUL),
            (14, WEB_EFFECT),
            (15, UTILIZATION),
        ):
            formula = '"' + ",".join(choices) + '"'
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            card.add_data_validation(dv)
            dv.add(card.cell(excel_row, 2))

        card.column_dimensions["A"].width = 28
        card.column_dimensions["B"].width = 100

    # Index sheet
    idx = wb.create_sheet("Index", 0)
    idx["A1"] = "case_id"
    idx["B1"] = "category"
    idx["C1"] = "request"
    idx["D1"] = "sheet"
    for i, row in enumerate(rows, 2):
        idx.cell(i, 1, row["case_id"])
        idx.cell(i, 2, row["category"])
        idx.cell(i, 3, row["request"])
        idx.cell(i, 4, row["case_id"])
    idx.column_dimensions["A"].width = 10
    idx.column_dimensions["B"].width = 8
    idx.column_dimensions["C"].width = 60
    idx.column_dimensions["D"].width = 10

    wb.save(path)


def write_readme(rows: list[dict], path: Path) -> None:
    n = len(rows)
    n0 = sum(1 for r in rows if (r.get("answer_without_web") or "").strip())
    ns = sum(1 for r in rows if (r.get("search_result_count") or 0) > 0)
    n1 = sum(1 for r in rows if (r.get("answer_with_web") or "").strip())
    path.write_text(
        "\n".join(
            [
                "# WEB検索効果レビューシート（再構成）",
                "",
                "**これはレビュー資料の再構成であり、Agent実装の変更ではない。**",
                "",
                "- heuristic / capability_route / Gate / Pipeline / Clarity / Stage3 / Stage4: **未変更**",
                "- 既存 `web_effect_review_hits.*` / `hits_batch_results/`: **削除していない**",
                "- 人間ラベルは空欄（レビュー結果は未作成）",
                "- 「どちらを採用するか」項目は設けない（差分観測）",
                "",
                "## 対象",
                "",
                f"- 件数: **{n}**（source: `web_effect_review_hits.json`）",
                f"- Webなし完成回答あり: **{n0}**",
                f"- 検索結果あり (count>0): **{ns}**",
                f"- Webあり完成回答あり: **{n1}**",
                "- 同一 `case_id` / `request` で without・search・with を対応付け",
                "",
                "## 1ケースの見方",
                "",
                "1. ユーザー要求（+ 質問の明確さ）",
                "2. Webなし完成回答（+ 十分性）",
                "3. Web検索結果 生データ（+ 有用性）",
                "4. Webあり完成回答（+ 変化 / 活用）",
                "5. 開発上のメモ",
                "",
                "## 記入ファイル",
                "",
                "| ファイル | 用途 |",
                "|----------|------|",
                "| `review_worksheet.xlsx` | **主推奨** Overview＋1ケース1シート、プルダウンあり |",
                "| `review_worksheet.csv` | 表形式バックアップ |",
                "| `COMPARE.md` | 読み比べカード |",
                "| `review_dataset.json` | 機械可読（検索JSON生データ含む） |",
                "",
                "## プルダウン",
                "",
                "| 列 | 選択肢 |",
                "|----|--------|",
                f"| question_clarity | {' / '.join(QUESTION_CLARITY)} |",
                f"| web_without_answer_quality | {' / '.join(WITHOUT_QUALITY)} |",
                f"| search_result_usefulness | {' / '.join(SEARCH_USEFUL)} |",
                f"| web_effect | {' / '.join(WEB_EFFECT)} |",
                f"| search_result_utilization | {' / '.join(UTILIZATION)} |",
                "",
                "詳細意味は Excel の `Choices` シート。",
                "",
                "## 再生成",
                "",
                "```text",
                "python analysis/build_web_effect_review_worksheet.py",
                "```",
                "",
            ]
        ),
        encoding="utf-8",
    )


def main() -> int:
    if not SOURCE.is_file():
        raise SystemExit(f"missing source: {SOURCE}")
    rows = load_rows()
    write_csv(rows, HERE / "review_worksheet.csv")
    write_json(rows, HERE / "review_dataset.json")
    write_compare_md(rows, HERE / "COMPARE.md")
    write_xlsx(rows, HERE / "review_worksheet.xlsx")
    write_readme(rows, HERE / "README.md")
    print(f"cases={len(rows)}")
    print("wrote review_worksheet.xlsx / .csv / review_dataset.json / COMPARE.md / README.md")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
