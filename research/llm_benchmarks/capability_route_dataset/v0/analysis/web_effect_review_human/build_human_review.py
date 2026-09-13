"""
人間向け Web効果レビュー資料の表示層を生成する。

- 既存原データ・heuristic / Gate / Pipeline / Stage3/4 は変更しない
- 検索結果は機械的整形のみ（LLM評価・要約禁止）
- 人間ラベルは空欄のまま
"""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SRC_DIR = Path(__file__).resolve().parents[1] / "web_effect_review"
SOURCE_JSON = SRC_DIR / "web_effect_review_hits.json"
BATCH_DIR = SRC_DIR / "hits_batch_results"
OUT_DIR = Path(__file__).resolve().parent

SNIPPET_MAX = 600

QUESTION_CLARITY = ["◎ 明確", "○ ほぼ明確", "△ 曖昧", "× 意味が破綻"]
WITHOUT_QUALITY = ["◎ 十分", "○ ほぼ十分", "△ 情報不足", "× 明確な問題あり"]
SEARCH_USEFUL = ["◎ とても役立つ", "○ 役立つ", "△ 一部役立つ", "× 役立たない", "― 検索失敗・比較不能"]
WITH_QUALITY = ["◎ 十分", "○ ほぼ十分", "△ 情報不足", "× 明確な問題あり"]
WEB_CHANGE = [
    "◎ 明確に良くなった",
    "○ 少し良くなった",
    "→ ほとんど変わらない",
    "△ 少し悪くなった",
    "× 明確に悪くなった",
    "― 比較不能",
]


def site_from_url(url: str | None) -> str:
    if not url:
        return ""
    try:
        host = urlparse(str(url)).netloc or ""
        return host[4:] if host.startswith("www.") else host
    except Exception:  # noqa: BLE001
        return ""


def clip_text(text: str, max_len: int = SNIPPET_MAX) -> tuple[str, bool]:
    s = str(text or "").strip()
    if len(s) <= max_len:
        return s, False
    return s[:max_len].rstrip() + "…", True


def human_readable_search_results(web_search: dict | None) -> dict:
    """既存 search result から機械的に表示用情報だけ抽出。LLM判断なし。"""
    ws = web_search or {}
    res = ws.get("result") or {}
    hits_raw = res.get("hits") if isinstance(res.get("hits"), list) else []
    query = ws.get("query") or res.get("query") or ""
    error = res.get("error")
    items = []
    any_truncated = False
    for i, hit in enumerate(hits_raw, 1):
        if not isinstance(hit, dict):
            items.append(
                {
                    "index": i,
                    "title": str(hit),
                    "site": "",
                    "url": "",
                    "content": "",
                    "content_truncated": False,
                }
            )
            continue
        title = str(hit.get("title") or "").strip()
        url = str(hit.get("url") or "").strip()
        snippet = hit.get("snippet")
        if snippet is None:
            snippet = hit.get("description")
        content_src = "" if snippet is None else str(snippet)
        content, truncated = clip_text(content_src, SNIPPET_MAX)
        if truncated:
            any_truncated = True
        items.append(
            {
                "index": i,
                "title": title,
                "site": site_from_url(url),
                "url": url,
                "content": content if content else "（内容テキストなし）",
                "content_truncated": truncated,
                "raw_fields_present": sorted(
                    k for k, v in hit.items() if v is not None and str(v).strip() != ""
                ),
            }
        )

    display_lines = [f"検索クエリ: {query}", ""]
    if not items:
        display_lines.append("検索結果は取得できませんでした。")
        if error:
            display_lines.append(f"（記録されたエラー: {error}）")
    else:
        for it in items:
            display_lines.append(f"検索結果 {it['index']}")
            display_lines.append(f"タイトル: {it['title'] or '（なし）'}")
            display_lines.append(f"サイト: {it['site'] or '（なし）'}")
            display_lines.append(f"URL: {it['url'] or '（なし）'}")
            display_lines.append("")
            display_lines.append("内容:")
            display_lines.append(it["content"])
            if it["content_truncated"]:
                display_lines.append("（検索結果本文は長いため省略。原データあり）")
            display_lines.append("")
    if any_truncated:
        display_lines.append("※一部の内容は表示用に省略。原JSONは変更していません。")

    return {
        "search_query": query,
        "result_count": res.get("hit_count")
        if res.get("hit_count") is not None
        else len(items),
        "error": error,
        "items": items,
        "display_text": "\n".join(display_lines).rstrip(),
        "raw_hits_preserved": True,
        "note": "機械的抽出のみ。有用性の自動判定は含まない。",
    }


def load_cases() -> list[dict]:
    data = json.loads(SOURCE_JSON.read_text(encoding="utf-8"))
    rows = []
    for case in data.get("cases") or []:
        case_id = str(case.get("case_id") or "")
        # 本バッチには Stage1 observation_id が無いため、追跡用IDを安定生成
        observation_id = f"webeffect-hits-{case_id}"
        source_batch = BATCH_DIR / f"{case_id}.json"
        human_search = human_readable_search_results(case.get("web_search"))
        raw_hits = ((case.get("web_search") or {}).get("result") or {}).get("hits") or []
        rows.append(
            {
                "case_id": case_id,
                "observation_id": observation_id,
                "category": case.get("band") or "",
                "request": case.get("request") or "",
                "answer_without_web": case.get("answer_without_web") or "",
                "answer_with_web": case.get("answer_with_web") or "",
                "human_readable_search_results": human_search,
                "search_results_raw": raw_hits,
                "source_ref": {
                    "review_hits_json": "web_effect_review/web_effect_review_hits.json",
                    "batch_case_json": (
                        f"web_effect_review/hits_batch_results/{case_id}.json"
                        if source_batch.is_file()
                        else None
                    ),
                },
                # 人間記入（空）
                "question_clarity": "",
                "question_fix_note": "",
                "web_without_answer_quality": "",
                "search_result_usefulness": "",
                "web_with_answer_quality": "",
                "web_search_effect": "",
                "human_note": "",
            }
        )
    return rows


def write_csv(rows: list[dict], path: Path) -> None:
    fields = [
        "case_id",
        "observation_id",
        "category",
        "request",
        "question_clarity",
        "question_fix_note",
        "answer_without_web",
        "web_without_answer_quality",
        "search_query",
        "search_result_count",
        "human_readable_search_results",
        "search_result_usefulness",
        "answer_with_web",
        "web_with_answer_quality",
        "web_search_effect",
        "human_note",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            hs = row["human_readable_search_results"]
            writer.writerow(
                {
                    "case_id": row["case_id"],
                    "observation_id": row["observation_id"],
                    "category": row["category"],
                    "request": row["request"],
                    "question_clarity": "",
                    "question_fix_note": "",
                    "answer_without_web": row["answer_without_web"],
                    "web_without_answer_quality": "",
                    "search_query": hs.get("search_query") or "",
                    "search_result_count": hs.get("result_count"),
                    "human_readable_search_results": hs.get("display_text") or "",
                    "search_result_usefulness": "",
                    "answer_with_web": row["answer_with_web"],
                    "web_with_answer_quality": "",
                    "web_search_effect": "",
                    "human_note": "",
                }
            )


def write_json(rows: list[dict], path: Path) -> None:
    payload = {
        "kind": "web_effect_human_review",
        "purpose": "Human review of web-search effect on final answers",
        "not_agent_gold": True,
        "no_llm_prejudgment": True,
        "no_auto_labels": True,
        "source": "web_effect_review/web_effect_review_hits.json",
        "case_count": len(rows),
        "counts": {
            "answer_without_web": sum(
                1 for r in rows if (r.get("answer_without_web") or "").strip()
            ),
            "search_results_present": sum(
                1
                for r in rows
                if (r.get("human_readable_search_results") or {}).get("result_count", 0)
                > 0
            ),
            "human_readable_search_results": sum(
                1
                for r in rows
                if (r.get("human_readable_search_results") or {}).get("display_text")
            ),
            "answer_with_web": sum(
                1 for r in rows if (r.get("answer_with_web") or "").strip()
            ),
        },
        "dropdowns": {
            "question_clarity": QUESTION_CLARITY,
            "web_without_answer_quality": WITHOUT_QUALITY,
            "search_result_usefulness": SEARCH_USEFUL,
            "web_with_answer_quality": WITH_QUALITY,
            "web_search_effect": WEB_CHANGE,
        },
        "cases": rows,
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_compare_md(rows: list[dict], path: Path) -> None:
    lines = [
        "# Web検索効果・人間レビュー比較カード",
        "",
        "プログラム内部仕様の理解は不要。完成回答と検索結果（人間向け表示）を読む。",
        "評価は「Web検索が必要だったか」ではなく、**検索を行った結果の変化**。",
        "",
    ]
    for row in rows:
        hs = row["human_readable_search_results"]
        lines.extend(
            [
                f"## Case {row['case_id']}",
                "",
                f"- observation_id: `{row['observation_id']}`",
                f"- category: `{row.get('category')}`",
                "",
                "### 質問",
                "",
                f"> {row.get('request')}",
                "",
                "**質問の意味は分かりますか？**（記入欄）",
                "",
                "---",
                "",
                "### Web検索なしの完成回答",
                "",
                "```text",
                row.get("answer_without_web") or "(empty)",
                "```",
                "",
                "**Web検索なしでも十分答えられていますか？**（記入欄）",
                "",
                "---",
                "",
                "### Web検索で得られた情報",
                "",
                "```text",
                hs.get("display_text") or "(empty)",
                "```",
                "",
                "**この検索結果は質問に答える材料として役立ちますか？**（記入欄）",
                "",
                "---",
                "",
                "### Web検索後の完成回答",
                "",
                "```text",
                row.get("answer_with_web") or "(empty)",
                "```",
                "",
                "**Web検索後の回答は十分ですか？**（記入欄）",
                "",
                "---",
                "",
                "### Web検索による変化",
                "",
                "**Web検索によって回答はどう変化しましたか？**（記入欄）",
                "",
                "### 人間メモ",
                "",
                "（自由記述）",
                "",
                "---",
                "",
            ]
        )
    path.write_text("\n".join(lines), encoding="utf-8")


def _dv(choices: list[str]) -> DataValidation:
    # Excel list はカンマ区切り。選択肢にカンマは含めないこと。
    formula = '"' + ",".join(choices) + '"'
    return DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)


def write_xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    label_fill = PatternFill("solid", fgColor="D6EAF8")
    wrap = Alignment(wrap_text=True, vertical="top")

    # Index
    idx = wb.active
    idx.title = "目次"
    idx["A1"] = "case_id"
    idx["B1"] = "observation_id"
    idx["C1"] = "質問"
    idx["D1"] = "シート"
    for cell in idx[1]:
        cell.fill = header_fill
        cell.font = header_font
    for i, row in enumerate(rows, 2):
        idx.cell(i, 1, row["case_id"])
        idx.cell(i, 2, row["observation_id"])
        idx.cell(i, 3, row["request"])
        idx.cell(i, 4, row["case_id"])
    idx.column_dimensions["A"].width = 10
    idx.column_dimensions["B"].width = 22
    idx.column_dimensions["C"].width = 55
    idx.column_dimensions["D"].width = 10

    # Overview
    ov = wb.create_sheet("一覧", 1)
    headers = [
        "case_id",
        "observation_id",
        "質問",
        "質問の明確さ",
        "質問の直し方メモ",
        "Webなし完成回答",
        "Webなし十分性",
        "検索クエリ",
        "検索結果（人間向け）",
        "検索結果の有用性",
        "Webあり完成回答",
        "Webあり十分性",
        "Web検索による変化",
        "人間メモ",
    ]
    for col, name in enumerate(headers, 1):
        cell = ov.cell(1, col, name)
        cell.fill = header_fill
        cell.font = header_font
    for i, row in enumerate(rows, 2):
        hs = row["human_readable_search_results"]
        vals = [
            row["case_id"],
            row["observation_id"],
            row["request"],
            "",
            "",
            row["answer_without_web"],
            "",
            hs.get("search_query") or "",
            hs.get("display_text") or "",
            "",
            row["answer_with_web"],
            "",
            "",
            "",
        ]
        for col, val in enumerate(vals, 1):
            cell = ov.cell(i, col, val)
            cell.alignment = wrap
        ov.row_dimensions[i].height = 140
    max_row = len(rows) + 1
    # D E F G I J L M N -> dropdowns on D,G,J,L,M
    for col, choices in (
        ("D", QUESTION_CLARITY),
        ("G", WITHOUT_QUALITY),
        ("J", SEARCH_USEFUL),
        ("L", WITH_QUALITY),
        ("M", WEB_CHANGE),
    ):
        dv = _dv(choices)
        ov.add_data_validation(dv)
        dv.add(f"{col}2:{col}{max_row}")
    widths = {
        "A": 10,
        "B": 22,
        "C": 28,
        "D": 14,
        "E": 18,
        "F": 40,
        "G": 14,
        "H": 20,
        "I": 40,
        "J": 16,
        "K": 40,
        "L": 14,
        "M": 18,
        "N": 24,
    }
    for k, w in widths.items():
        ov.column_dimensions[k].width = w
    ov.freeze_panes = "C2"
    ov.auto_filter.ref = f"A1:N{max_row}"

    # Choices
    ch = wb.create_sheet("選択肢の説明", 2)
    ch["A1"] = "項目"
    ch["B1"] = "選択肢"
    ch["C1"] = "意味"
    meanings = {
        "質問の明確さ": {
            "◎ 明確": "意味が分かる",
            "○ ほぼ明確": "だいたい分かる",
            "△ 曖昧": "解釈が分かれる",
            "× 意味が破綻": "質問として成立しにくい（元文は変えない）",
        },
        "Webなし十分性": {
            "◎ 十分": "この回答だけで足りる",
            "○ ほぼ十分": "大きな欠落はない",
            "△ 情報不足": "足りない",
            "× 明確な問題あり": "誤り・空疎など明確な問題",
        },
        "検索結果の有用性": {
            "◎ とても役立つ": "回答材料として非常に使える",
            "○ 役立つ": "使える",
            "△ 一部役立つ": "一部だけ",
            "× 役立たない": "使えない",
            "― 検索失敗・比較不能": "結果なし等で比較できない",
        },
        "Webあり十分性": {
            "◎ 十分": "この回答だけで足りる",
            "○ ほぼ十分": "大きな欠落はない",
            "△ 情報不足": "足りない",
            "× 明確な問題あり": "明確な問題",
        },
        "Web検索による変化": {
            "◎ 明確に良くなった": "明確な改善",
            "○ 少し良くなった": "改善はあるが小さい",
            "→ ほとんど変わらない": "実質差が小さい",
            "△ 少し悪くなった": "やや悪化",
            "× 明確に悪くなった": "明確な悪化",
            "― 比較不能": "比較できない",
        },
    }
    r = 2
    mapping = [
        ("質問の明確さ", QUESTION_CLARITY),
        ("Webなし十分性", WITHOUT_QUALITY),
        ("検索結果の有用性", SEARCH_USEFUL),
        ("Webあり十分性", WITH_QUALITY),
        ("Web検索による変化", WEB_CHANGE),
    ]
    for field, opts in mapping:
        for opt in opts:
            ch.cell(r, 1, field)
            ch.cell(r, 2, opt)
            ch.cell(r, 3, meanings[field].get(opt, ""))
            r += 1
    ch.column_dimensions["A"].width = 20
    ch.column_dimensions["B"].width = 22
    ch.column_dimensions["C"].width = 40

    # Per-case cards
    for row in rows:
        name = str(row["case_id"])[:31]
        card = wb.create_sheet(name)
        card["A1"] = "項目"
        card["B1"] = "内容"
        card["A1"].fill = header_fill
        card["B1"].fill = header_fill
        card["A1"].font = header_font
        card["B1"].font = header_font
        hs = row["human_readable_search_results"]
        blocks = [
            ("case_id", row["case_id"]),
            ("observation_id", row["observation_id"]),
            ("質問", row["request"]),
            ("質問の意味は分かりますか？（記入）", ""),
            ("どう直せば分かりやすいか（任意）", ""),
            ("Web検索なしの完成回答", row["answer_without_web"]),
            ("Web検索なしでも十分ですか？（記入）", ""),
            ("検索クエリ", hs.get("search_query") or ""),
            ("Web検索で得られた情報", hs.get("display_text") or ""),
            ("検索結果は役立ちますか？（記入）", ""),
            ("Web検索後の完成回答", row["answer_with_web"]),
            ("Web検索後の回答は十分ですか？（記入）", ""),
            ("Web検索によって回答はどう変化しましたか？（記入）", ""),
            ("人間メモ（自由記述）", ""),
        ]
        for i, (label, value) in enumerate(blocks, 2):
            a = card.cell(i, 1, label)
            b = card.cell(i, 2, value)
            a.fill = label_fill
            a.alignment = wrap
            b.alignment = wrap
            card.row_dimensions[i].height = 28 if "記入" in label or "任意" in label or "メモ" in label else (
                110 if len(str(value)) > 60 else 40
            )
        # dropdown rows: 5,8,11,13,14
        for excel_row, choices in (
            (5, QUESTION_CLARITY),
            (8, WITHOUT_QUALITY),
            (11, SEARCH_USEFUL),
            (13, WITH_QUALITY),
            (14, WEB_CHANGE),
        ):
            dv = _dv(choices)
            card.add_data_validation(dv)
            dv.add(card.cell(excel_row, 2))
        card.column_dimensions["A"].width = 36
        card.column_dimensions["B"].width = 100

    wb.save(path)


def write_readme(rows: list[dict], path: Path) -> None:
    n = len(rows)
    c = {
        "without": sum(1 for r in rows if (r.get("answer_without_web") or "").strip()),
        "search": sum(
            1
            for r in rows
            if (r.get("human_readable_search_results") or {}).get("result_count", 0) > 0
        ),
        "human": sum(
            1
            for r in rows
            if (r.get("human_readable_search_results") or {}).get("display_text")
        ),
        "with": sum(1 for r in rows if (r.get("answer_with_web") or "").strip()),
    }
    path.write_text(
        "\n".join(
            [
                "# Web検索効果・人間レビュー資料",
                "",
                "> この資料は、人間がWeb検索の効果を評価するためのレビュー資料である。",
                "",
                "> 原検索結果をLLMが評価・要約したものではなく、可能な限り元検索結果から機械的に表示用情報を抽出している。",
                "",
                "> 評価対象は「Web検索が必要だったか」ではなく、実際にWeb検索を行った場合に完成回答がどのように変化したかである。",
                "",
                "> Web検索なしの回答、検索結果、Web検索後の回答の3者を比較する。",
                "",
                "> 人間レビュー結果を正解ラベルとして自動利用しない。",
                "",
                "## 注意",
                "",
                "- heuristic / Gate / Pipeline / Clarity / Stage3 / Stage4 / 既存原データ: **未変更**",
                "- 元データ: `../web_effect_review/web_effect_review_hits.json`（削除・上書きしない）",
                "- 人間記入欄は空欄（レビュー結果は未作成）",
                "",
                "## 対象件数",
                "",
                f"- ケース数: **{n}**",
                f"- human-readable 検索結果: **{c['human']}**",
                f"- Webなし完成回答: **{c['without']}**",
                f"- 検索結果あり: **{c['search']}**",
                f"- Webあり完成回答: **{c['with']}**",
                "- 追跡: 各行に `case_id` と `observation_id`（`webeffect-hits-<case_id>`）",
                "",
                "## 使い方",
                "",
                "1. `review_worksheet.xlsx` を開く（推奨）",
                "2. `目次` からケースシートへ、または `一覧` で横断記入",
                "3. プルダウンで評価し、必要ならメモを書く",
                "",
                "## プルダウン項目",
                "",
                "| 項目 | 選択肢 |",
                "|------|--------|",
                f"| 質問の明確さ | {' / '.join(QUESTION_CLARITY)} |",
                f"| Webなし十分性 | {' / '.join(WITHOUT_QUALITY)} |",
                f"| 検索結果の有用性 | {' / '.join(SEARCH_USEFUL)} |",
                f"| Webあり十分性 | {' / '.join(WITH_QUALITY)} |",
                f"| Web検索による変化 | {' / '.join(WEB_CHANGE)} |",
                "",
                "## ファイル",
                "",
                "| ファイル | 内容 |",
                "|----------|------|",
                "| `review_worksheet.xlsx` | 1ケース1カード＋一覧＋プルダウン |",
                "| `review_worksheet.csv` | 表形式 |",
                "| `review_dataset.json` | 機械可読（原hitsも同梱） |",
                "| `COMPARE.md` | Markdown比較カード |",
                "",
            ]
        ),
        encoding="utf-8",
    )


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if not SOURCE_JSON.is_file():
        raise SystemExit(f"missing source: {SOURCE_JSON}")
    rows = load_cases()
    write_csv(rows, OUT_DIR / "review_worksheet.csv")
    write_json(rows, OUT_DIR / "review_dataset.json")
    write_compare_md(rows, OUT_DIR / "COMPARE.md")
    write_xlsx(rows, OUT_DIR / "review_worksheet.xlsx")
    write_readme(rows, OUT_DIR / "README.md")
    print(f"wrote {OUT_DIR}")
    print(f"cases={len(rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
